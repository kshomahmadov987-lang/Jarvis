"""
Office Controller for Jarvis AI
Handles document creation and manipulation (Word, Excel)
"""

import os
from typing import Optional, List, Dict
from datetime import datetime


class OfficeController:
    """
    Manages Microsoft Office documents.
    Supports Word (.docx) and Excel (.xlsx) files.
    """
    
    def __init__(self, default_dir: str = None):
        """
        Initialize office controller.
        
        Args:
            default_dir: Default directory for saving documents
        """
        self.default_dir = default_dir or os.path.join(os.getcwd(), "jarvis_documents")
        self._ensure_default_dir()
        self._docx_available = False
        self._openpyxl_available = False
        self._initialize_libraries()
        
    def _ensure_default_dir(self):
        """Create default directory if it doesn't exist."""
        if not os.path.exists(self.default_dir):
            os.makedirs(self.default_dir)
            print(f"[OFFICE] Created documents directory: {self.default_dir}")
            
    def _initialize_libraries(self):
        """Initialize document processing libraries."""
        try:
            from docx import Document
            self._docx_available = True
            print("[OFFICE] python-docx initialized")
        except ImportError:
            print("[WARNING] python-docx not installed. Word support disabled.")
            
        try:
            from openpyxl import Workbook
            self._openpyxl_available = True
            print("[OFFICE] openpyxl initialized")
        except ImportError:
            print("[WARNING] openpyxl not installed. Excel support disabled.")
            
    # Word Document Methods
    def create_word_document(self, filename: str, content: str = "", title: str = None) -> Optional[str]:
        """
        Create a new Word document.
        
        Args:
            filename: Name of the file (without extension)
            content: Initial content text
            title: Optional document title
            
        Returns:
            Full path to created file or None if failed
        """
        if not self._docx_available:
            print("[ERROR] Word support not available")
            return None
            
        try:
            from docx import Document
            
            doc = Document()
            
            if title:
                doc.add_heading(title, 0)
                
            if content:
                doc.add_paragraph(content)
                
            # Add timestamp
            doc.add_paragraph(f"\nCreated by Jarvis AI on {datetime.now().strftime('%Y-%m-%d %H:%M')}")
            
            filepath = os.path.join(self.default_dir, f"{filename}.docx")
            doc.save(filepath)
            
            print(f"[OFFICE] Created Word document: {filepath}")
            return filepath
            
        except Exception as e:
            print(f"[ERROR] Failed to create Word document: {e}")
            return None
            
    def append_to_word(self, filename: str, text: str, heading: bool = False) -> bool:
        """
        Append text to an existing Word document.
        
        Args:
            filename: Name of the file (without extension)
            text: Text to append
            heading: If True, add as heading
            
        Returns:
            True if successful, False otherwise
        """
        if not self._docx_available:
            return False
            
        try:
            from docx import Document
            
            filepath = os.path.join(self.default_dir, f"{filename}.docx")
            
            if not os.path.exists(filepath):
                print(f"[ERROR] File not found: {filepath}")
                return False
                
            doc = Document(filepath)
            
            if heading:
                doc.add_heading(text, level=1)
            else:
                doc.add_paragraph(text)
                
            doc.save(filepath)
            print(f"[OFFICE] Appended to {filepath}")
            return True
            
        except Exception as e:
            print(f"[ERROR] Failed to append to Word document: {e}")
            return False
            
    # Excel Spreadsheet Methods
    def create_excel_sheet(self, filename: str, data: List[List], headers: List[str] = None) -> Optional[str]:
        """
        Create a new Excel spreadsheet.
        
        Args:
            filename: Name of the file (without extension)
            data: 2D list of cell values
            headers: Optional header row
            
        Returns:
            Full path to created file or None if failed
        """
        if not self._openpyxl_available:
            print("[ERROR] Excel support not available")
            return None
            
        try:
            from openpyxl import Workbook
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Jarvis Sheet"
            
            row_num = 1
            
            # Add headers
            if headers:
                for col_num, header in enumerate(headers, 1):
                    ws.cell(row=row_num, column=col_num, value=header)
                row_num += 1
                
            # Add data
            for row_data in data:
                for col_num, value in enumerate(row_data, 1):
                    ws.cell(row=row_num, column=col_num, value=value)
                row_num += 1
                
            filepath = os.path.join(self.default_dir, f"{filename}.xlsx")
            wb.save(filepath)
            
            print(f"[OFFICE] Created Excel sheet: {filepath}")
            return filepath
            
        except Exception as e:
            print(f"[ERROR] Failed to create Excel sheet: {e}")
            return None
            
    def read_excel_sheet(self, filename: str) -> Optional[List[List]]:
        """
        Read data from an Excel spreadsheet.
        
        Args:
            filename: Name of the file (without extension)
            
        Returns:
            2D list of cell values or None if failed
        """
        if not self._openpyxl_available:
            return None
            
        try:
            from openpyxl import load_workbook
            
            filepath = os.path.join(self.default_dir, f"{filename}.xlsx")
            
            if not os.path.exists(filepath):
                print(f"[ERROR] File not found: {filepath}")
                return None
                
            wb = load_workbook(filename=filepath, data_only=True)
            ws = wb.active
            
            data = []
            for row in ws.iter_rows(values_only=True):
                data.append(list(row))
                
            return data
            
        except Exception as e:
            print(f"[ERROR] Failed to read Excel sheet: {e}")
            return None
            
    def list_documents(self) -> Dict[str, List[str]]:
        """
        List all documents in the default directory.
        
        Returns:
            Dictionary with 'word' and 'excel' lists of filenames
        """
        result = {'word': [], 'excel': []}
        
        if not os.path.exists(self.default_dir):
            return result
            
        for filename in os.listdir(self.default_dir):
            if filename.endswith('.docx'):
                result['word'].append(filename)
            elif filename.endswith('.xlsx'):
                result['excel'].append(filename)
                
        return result
